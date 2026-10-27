"""Monta a planilha de relatório: melhores ofertas, comparativo, mudanças e histórico."""
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

AZUL = PatternFill("solid", fgColor="1F4E78")
VERDE = PatternFill("solid", fgColor="DEF7E5")
VERMELHO = PatternFill("solid", fgColor="FDE2E2")
AMARELO = PatternFill("solid", fgColor="FFF4CE")
CINZA = PatternFill("solid", fgColor="F2F2F2")
FONTE = "Calibri"
REAIS = '"R$" #,##0.00'
PORCENTO = "0.0%"
CENTRO = Alignment(horizontal="center")


def titulo(aba, texto, subtitulo):
    aba["A1"] = texto
    aba["A1"].font = Font(name=FONTE, bold=True, size=14, color="1F4E78")
    aba["A2"] = subtitulo
    aba["A2"].font = Font(name=FONTE, italic=True, color="666666")


def cabecalho(aba, colunas, linha=4):
    """colunas = [(nome, largura), ...]"""
    for numero, (nome, largura) in enumerate(colunas, start=1):
        celula = aba.cell(linha, numero, nome)
        celula.fill = AZUL
        celula.font = Font(name=FONTE, bold=True, color="FFFFFF")
        celula.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        aba.column_dimensions[celula.column_letter].width = largura
    aba.freeze_panes = aba.cell(linha + 1, 1)


def escrever(aba, linha, valores, formatos=None, cor=None):
    formatos = formatos or {}
    for numero, valor in enumerate(valores, start=1):
        celula = aba.cell(linha, numero, valor)
        celula.font = Font(name=FONTE)
        if numero in formatos:
            celula.number_format = formatos[numero]
        if numero > 2:
            celula.alignment = CENTRO
        if cor:
            celula.fill = cor


def para_impressao(aba):
    aba.page_setup.orientation = "landscape"
    aba.sheet_properties.pageSetUpPr.fitToPage = True
    aba.page_setup.fitToWidth, aba.page_setup.fitToHeight = 1, 0


def melhores_ofertas(aba, produtos, lojas):
    cabecalho(aba, [("Produto", 40), ("EAN", 16), ("Melhor oferta", 24), ("Melhor preço", 14),
                    ("Mais caro em", 24), ("Maior preço", 14), ("Economia (R$)", 14),
                    ("Economia (%)", 13), ("Lojas com estoque", 12)])
    linha = 5
    for (ean, nome), por_loja in produtos.items():
        ofertas = sorted((p["preco"], loja) for loja, p in por_loja.items() if p["disponivel"])
        if not ofertas:
            escrever(aba, linha, [nome, ean, "Indisponível em todas"], cor=VERMELHO)
            for coluna in range(4, 10):
                aba.cell(linha, coluna).fill = VERMELHO
        else:
            (menor, loja_menor), (maior, loja_maior) = ofertas[0], ofertas[-1]
            uma_so = len(ofertas) == 1 or menor == maior
            escrever(aba, linha, [
                nome, ean, loja_menor, menor,
                None if uma_so else loja_maior, None if uma_so else maior,
                None if uma_so else round(maior - menor, 2),
                None if uma_so else (maior - menor) / maior,
                f"{len(ofertas)} de {len(lojas)}"],
                formatos={4: REAIS, 6: REAIS, 7: REAIS, 8: PORCENTO})
            aba.cell(linha, 4).fill = VERDE
            aba.cell(linha, 4).font = Font(name=FONTE, bold=True, color="1D6B35")
            if uma_so:
                aviso = "Só esta loja tem estoque" if len(ofertas) == 1 else "Mesmo preço nas lojas"
                aba.cell(linha, 5, aviso).font = Font(name=FONTE, italic=True, color="666666")
        linha += 1


def comparativo(aba, produtos, lojas, lojas_com_falha):
    cabecalho(aba, [("Produto", 40), ("EAN", 16)] + [(loja, 22) for loja in lojas])
    linha = 5
    for (ean, nome), por_loja in produtos.items():
        escrever(aba, linha, [nome, ean])
        disponiveis = [p["preco"] for p in por_loja.values() if p["disponivel"]]
        menor = min(disponiveis) if disponiveis else None
        for numero, loja in enumerate(lojas, start=3):
            celula = aba.cell(linha, numero)
            celula.alignment = CENTRO
            p = por_loja.get(loja)
            if loja in lojas_com_falha:
                celula.value, celula.fill = "Sem coleta", AMARELO
                celula.font = Font(name=FONTE, italic=True)
            elif p is None:
                celula.value, celula.fill = "Não vende", CINZA
                celula.font = Font(name=FONTE, italic=True, color="999999")
            elif not p["disponivel"]:
                celula.value, celula.fill = "Indisponível", CINZA
                celula.font = Font(name=FONTE, italic=True, color="999999")
            else:
                celula.value, celula.number_format = p["preco"], REAIS
                melhor = p["preco"] == menor
                celula.font = Font(name=FONTE, bold=melhor, color="1D6B35" if melhor else None)
                if melhor:
                    celula.fill = VERDE
        linha += 1
    aba.cell(linha + 1, 1, "Verde = melhor preço entre as lojas com estoque.").font = Font(
        name=FONTE, italic=True, color="666666")


def o_que_mudou(aba, mudancas):
    cabecalho(aba, [("Loja", 24), ("Produto", 40), ("Antes", 13), ("Agora", 13),
                    ("Diferença (R$)", 14), ("Diferença (%)", 13), ("Situação", 20)])
    if not mudancas:
        aba["A5"] = ("Nada mudou desde a coleta anterior (ou esta é a primeira coleta: "
                     "a partir da próxima, as mudanças aparecem aqui).")
        aba["A5"].font = Font(name=FONTE, italic=True, color="666666")
        return
    cores = {"Subiu": VERMELHO, "Baixou": VERDE}
    ordem = ["Baixou", "Subiu", "Ficou indisponível", "Voltou ao estoque", "Produto novo",
             "Saiu do site"]
    mudancas = sorted(mudancas, key=lambda m: (ordem.index(m["situacao"]), m["produto"]))
    for linha, m in enumerate(mudancas, start=5):
        diferenca = round(m["agora"] - m["antes"], 2) if m["antes"] and m["agora"] else None
        escrever(aba, linha, [
            m["loja"], m["produto"], m["antes"], m["agora"], diferenca,
            diferenca / m["antes"] if diferenca is not None else None,
            ("▲ " if m["situacao"] == "Subiu" else "▼ " if m["situacao"] == "Baixou" else "")
            + m["situacao"]],
            formatos={3: REAIS, 4: REAIS, 5: REAIS, 6: PORCENTO},
            cor=cores.get(m["situacao"], AMARELO))


def evolucao(aba, linhas_do_historico):
    datas = sorted({h["data_hora"] for h in linhas_do_historico})
    cabecalho(aba, [("Produto", 40), ("Loja", 24)] +
              [(datetime.fromisoformat(d).strftime("%d/%m/%Y\n%H:%M"), 13) for d in datas])
    aba.row_dimensions[4].height = 30
    series = {}
    for h in linhas_do_historico:
        valor = h["preco"] if h["disponivel"] else "Indisp."
        series.setdefault((h["produto"], h["loja"]), {})[h["data_hora"]] = valor
    for linha, ((produto, loja), valores) in enumerate(series.items(), start=5):
        escrever(aba, linha, [produto, loja] + [valores.get(d) for d in datas],
                 formatos={c: REAIS for c in range(3, len(datas) + 3)})


def ocorrencias(aba, falhas):
    cabecalho(aba, [("Loja", 24), ("O que aconteceu", 55), ("Print da tela", 22)])
    for linha, f in enumerate(falhas, start=5):
        escrever(aba, linha, [f["loja"], f["motivo"].capitalize(), "Abrir print"], cor=VERMELHO)
        link = aba.cell(linha, 3)
        link.hyperlink = f["print"]
        link.font = Font(name=FONTE, color="0563C1", underline="single")


def gerar_relatorio(caminho, quando, lojas, precos, mudancas, linhas_do_historico, falhas):
    produtos = {}
    for p in sorted(precos, key=lambda p: p["produto"]):
        produtos.setdefault((p["ean"], p["produto"]), {})[p["loja"]] = p
    lojas_com_falha = {f["loja"] for f in falhas}
    coleta = f"Coleta de {quando:%d/%m/%Y às %H:%M}"
    if falhas:
        coleta += f"  ·  atenção: {len(falhas)} loja(s) sem coleta (ver aba Ocorrências)"

    livro = Workbook()
    abas = [
        ("Melhores ofertas", "Melhor oferta de cada produto", melhores_ofertas, (produtos, lojas)),
        ("Comparativo", "Preço de cada produto em cada loja", comparativo,
         (produtos, lojas, lojas_com_falha)),
        ("O que mudou", "Mudanças desde a coleta anterior de cada loja", o_que_mudou, (mudancas,)),
        ("Histórico", "Preço de cada produto em todas as coletas", evolucao,
         (linhas_do_historico,)),
    ]
    if falhas:
        abas.append(("Ocorrências", "Lojas que não puderam ser coletadas", ocorrencias, (falhas,)))
    for numero, (nome, texto, montar, dados) in enumerate(abas):
        aba = livro.active if numero == 0 else livro.create_sheet()
        aba.title = nome
        titulo(aba, texto, coleta)
        montar(aba, *dados)
        para_impressao(aba)
    livro.save(caminho)
